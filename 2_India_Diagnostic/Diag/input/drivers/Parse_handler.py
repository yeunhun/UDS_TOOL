import csv
from collections import defaultdict
from pathlib import Path

grouped_cases = defaultdict(list)

# Canonical column names we expect in the Excel sheet (header row)
_EXCEL_COL_ALIASES = {
    "tc_id/pc_id": "tc_id",
    "tc_id": "tc_id",
    "testcase_description": "step_desc",
    "testcase description": "step_desc",
    "service_id": "service_id",
    "subservice_id(did)": "subfunc",
    "subservice_id": "subfunc",
    "expected_response_data": "expected",
    "write_data": "write_data",
    "addressing": "addressing",
    "status_mask": "status_mask",
    "communication_type": "communication_type",
    "controltype": "controltype",
    # optional
    "format_type": "format_type",
}

def _norm_header(h):
    return str(h).strip().lower().replace("\n", " ").replace("\r", " ")

def _as_str(v):
    if v is None:
        return ""
    return str(v).strip()

def load_testcases(file_path):
    """
    Returns: dict[str, list[tuple]]
      grouped_cases[tc_id] -> list of steps
      step tuple layout must match what uds_client.py expects:
        (tc_id, step_desc, service_id, subfunc_or_did, expected_response,
         write_data, addressing, format_type, status_mask, communication_type, controltype)
    """
    grouped_cases.clear()
    p = Path(file_path)

    try:
        if p.suffix.lower() in (".xlsx", ".xlsm"):
            return _load_testcases_from_excel(p)
        else:
            return _load_testcases_from_csv_like(p)
    except Exception as e:
        print(f"Error parsing testcases: {e}")
        return {}

def _load_testcases_from_csv_like(p: Path):
    with p.open("r", newline="") as f:
        reader = csv.reader(f)
        _ = next(reader, None)  # header

        for row in reader:
            if not row or len(row) < 5:
                continue

            tc_id = (row[0] or "").strip()
            if not tc_id:
                continue

            step_desc = (row[1] or "").strip()
            service_id = (row[2] or "").strip()
            subfunction_or_did = (row[3] or "").strip()
            expected_response = (row[4] or "").strip()

            write_data = (row[5] or "").strip() if len(row) > 5 else ""
            addressing = (row[6] or "physical").strip().lower() if len(row) > 6 else "physical"
            format_type = (row[7] or "hex").strip() if len(row) > 7 else "hex"
            status_mask = (row[8] or "").strip() if len(row) > 8 else ""
            communication_type = (row[9] or "").strip() if len(row) > 9 else ""
            controltype = (row[10] or "").strip() if len(row) > 10 else ""

            grouped_cases[tc_id].append((
                tc_id,
                step_desc,
                service_id,
                subfunction_or_did,
                expected_response,
                write_data,
                addressing,
                format_type,
                status_mask,
                communication_type,
                controltype
            ))

    return grouped_cases

def _load_testcases_from_excel(p: Path):
    from openpyxl import load_workbook

    wb = load_workbook(p, data_only=True)
    ws = wb[wb.sheetnames[0]]  # default: first sheet

    header_row = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col_map = {}
    for idx, h in enumerate(header_row, start=1):
        key = _norm_header(h)
        if not key:
            continue
        if key in _EXCEL_COL_ALIASES:
            col_map[_EXCEL_COL_ALIASES[key]] = idx

    required = ["tc_id", "step_desc", "service_id", "subfunc", "expected", "write_data", "addressing"]
    missing = [k for k in required if k not in col_map]
    if missing:
        raise ValueError(
            f"Excel header missing required columns: {missing}. "
            f"Found headers: {[str(h) for h in header_row if h is not None]}"
        )

    idx_format = col_map.get("format_type")
    idx_status_mask = col_map.get("status_mask")
    idx_comm_type = col_map.get("communication_type")
    idx_control = col_map.get("controltype")

    for r in range(2, ws.max_row + 1):
        tc_id = _as_str(ws.cell(r, col_map["tc_id"]).value)
        if not tc_id:
            continue

        step_desc = _as_str(ws.cell(r, col_map["step_desc"]).value)
        service_id = _as_str(ws.cell(r, col_map["service_id"]).value)
        subfunc = _as_str(ws.cell(r, col_map["subfunc"]).value)
        expected = _as_str(ws.cell(r, col_map["expected"]).value)
        write_data = _as_str(ws.cell(r, col_map["write_data"]).value)
        addressing = _as_str(ws.cell(r, col_map["addressing"]).value).lower() or "physical"

        format_type = _as_str(ws.cell(r, idx_format).value) if idx_format else "hex"
        status_mask = _as_str(ws.cell(r, idx_status_mask).value) if idx_status_mask else ""
        communication_type = _as_str(ws.cell(r, idx_comm_type).value) if idx_comm_type else ""
        controltype = _as_str(ws.cell(r, idx_control).value) if idx_control else ""

        grouped_cases[tc_id].append((
            tc_id,
            step_desc,
            service_id,
            subfunc,
            expected,
            write_data,
            addressing,
            format_type or "hex",
            status_mask,
            communication_type,
            controltype
        ))

    return grouped_cases

def export_grouped_cases_to_txt(grouped_cases_dict, out_path):
    """
    Optional helper to keep legacy components (e.g., HTML report generator) working.
    Creates the same CSV-like testcase.txt that the old pipeline used.
    """
    out_p = Path(out_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)

    header = [
        "TC_ID", "Step_Description", "Service_ID", "SubService_ID(DID)",
        "Expected_Response_Data", "Write_Data", "Addressing", "Format_Type",
        "Status_Mask", "Communication_Type", "controltype"
    ]

    with out_p.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        for tc_id, steps in grouped_cases_dict.items():
            for step in steps:
                w.writerow(list(step))
    return str(out_p)
